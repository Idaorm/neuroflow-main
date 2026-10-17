"""Regression test for the EEG Excel export on Windows."""

import importlib.util
import os
import sys
import tempfile
import threading
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook


SERVER_DIR = Path(__file__).resolve().parents[1] / "emotiv" / "cortex-example" / "python"


class ExcelExportTest(unittest.TestCase):
    def test_creates_workbook_with_recorded_data(self):
        sys.path.insert(0, str(SERVER_DIR))
        try:
            spec = importlib.util.spec_from_file_location("emotiv_server_test", SERVER_DIR / "server.py")
            module = importlib.util.module_from_spec(spec)
            with mock.patch.dict(os.environ, {"EMOTIV_CLIENT_ID": "test", "EMOTIV_CLIENT_SECRET": "test"}):
                # Importing server.py normally starts the hardware connection.
                with mock.patch.object(threading.Thread, "start"):
                    spec.loader.exec_module(module)

            collector = module.collector
            collector.labels = {"met": ["focus", "engagement"],
                                "mot": ["Q0", "ACCX"], "pow": ["AF3/alpha", "AF3/theta"]}
            collector.current_marker = "inizio slide1"
            collector.on_new_met_data(data={"time": 1234.5, "met": [0.4123456789, 0.876543219]})
            collector.on_new_mot_data(data={"time": 1234.6, "mot": [0.987654321, -0.123456789]})
            collector.on_new_pow_data(data={"time": 1234.7, "pow": [1.234567891, 0.0000123456789]})
            collector.marker_buffer = [{"time": 1234.5, "label": "inizio slide1", "samples_at_set": 1}]
            self.assertTrue(collector.set_stop_marker(1240.75))
            self.assertFalse(collector.set_stop_marker(1241))

            with tempfile.TemporaryDirectory() as folder:
                filename = str(Path(folder) / "session.xlsx")
                self.assertEqual(collector.save_to_excel(filename, "Test"), filename)
                workbook = load_workbook(filename, read_only=True)
                try:
                    self.assertIn("Performance", workbook.sheetnames)
                    self.assertIn("Markers", workbook.sheetnames)
                    for sheet, time_value, values in [
                        ("Performance", 1234.5, [0.4123456789, 0.876543219]),
                        ("Motion", 1234.6, [0.987654321, -0.123456789]),
                        ("BandPower", 1234.7, [1.234567891, 0.0000123456789]),
                    ]:
                        self.assertEqual(workbook[sheet]["A2"].value, time_value)
                        self.assertEqual(workbook[sheet]["B2"].value, values[0])
                        self.assertEqual(workbook[sheet]["C2"].value, values[1])
                    self.assertEqual(workbook["Performance"]["D2"].value, "inizio slide1")
                    self.assertEqual(workbook["Markers"]["B2"].value, "inizio slide1")
                    self.assertEqual(workbook["Markers"]["A2"].value, 1234.5)
                    self.assertEqual(workbook["Markers"]["B3"].value, "stop")
                    self.assertEqual(workbook["Markers"]["A3"].value, 1240.75)
                    self.assertEqual(workbook["Markers"].max_row, 3)
                finally:
                    workbook.close()
        finally:
            sys.path.remove(str(SERVER_DIR))


if __name__ == "__main__":
    unittest.main()
